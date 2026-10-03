from pathlib import Path
import hashlib
import json
import subprocess
import threading
import wave
import zipfile
import pytest
from PIL import Image
import fitz
from openpyxl import Workbook,load_workbook
from docx import Document
from pptx import Presentation
from any2any.engine import ConversionEngine,normalize_extension
from any2any.model import Settings,ConversionError,Cancelled
from any2any.tools import locate


@pytest.fixture
def engine(): return ConversionEngine()


@pytest.fixture
def sources(tmp_path):
    image=tmp_path/'image.jpg'; Image.new('RGB',(96,64),(60,130,210)).save(image)
    png=tmp_path/'image.png'; Image.new('RGB',(96,64),(60,130,210)).save(png)
    webp=tmp_path/'image.webp'; Image.new('RGB',(96,64),'red').save(webp)
    text=tmp_path/'note.txt'; text.write_text('Hello Any2Any\nA meaningful interpretation.',encoding='utf-8')
    xlsx=tmp_path/'book.xlsx'; book=Workbook(); book.active.append(['x','y']); book.active.append([1,2]); book.active.append([3,5]); book.save(xlsx)
    docx=tmp_path/'document.docx'; doc=Document(); doc.add_paragraph('A document with useful text'); doc.save(docx)
    pptx=tmp_path/'slides.pptx'; presentation=Presentation(); slide=presentation.slides.add_slide(presentation.slide_layouts[1]); slide.shapes.title.text='Slide title'; slide.placeholders[1].text='Content'; presentation.save(pptx)
    pdf=tmp_path/'document.pdf'; document=fitz.open()
    for title in ('Page one','Page two'):
        page=document.new_page(width=300,height=200); page.insert_text((30,60),title)
    document.save(pdf); document.close()
    binary=tmp_path/'unknown.bin'; binary.write_bytes(bytes(range(256))*8)
    wav=tmp_path/'sound.wav'
    import math,struct
    with wave.open(str(wav),'wb') as audio:
        audio.setparams((1,2,22050,0,'NONE','not compressed'))
        audio.writeframes(b''.join(struct.pack('<h',int(4000*math.sin(i*2*math.pi*440/22050))) for i in range(22050)))
    archive=tmp_path/'files.zip'
    with zipfile.ZipFile(archive,'w') as zip_:
        zip_.writestr('notes/hello.txt','Archive text'); zip_.writestr('../unsafe.txt','DO NOT EXTRACT')
    return locals()


def test_signature_over_extension(engine,sources,tmp_path):
    disguised=tmp_path/'fake.mp3'; disguised.write_bytes(sources['png'].read_bytes())
    info=engine.detector.detect(disguised)
    assert info.format=='png' and info.category=='image' and info.evidence=='signature'
    for key,fmt in [('xlsx','xlsx'),('docx','docx'),('pptx','pptx'),('pdf','pdf'),('wav','wav'),('archive','zip')]:
        assert engine.detector.detect(sources[key]).format==fmt


@pytest.mark.parametrize('source,target,label,contains',[
    ('image','png','Native conversion','Pillow'),('webp','jpg','Native conversion','Pillow'),
    ('png','wav','Semantic conversion','image sonification'),('xlsx','mp3','Semantic conversion','spreadsheet'),
    ('pdf','mp4','Semantic conversion','PDF page renderer'),('text','png','Semantic conversion','wrapped text'),
    ('archive','txt','Semantic conversion','safe archive'),('binary','xyz','Binary interpretation','envelope'),
])
def test_route_selection(engine,sources,source,target,label,contains):
    info,plan=engine.preview(sources[source],target)
    assert plan.label==label and contains in plan.route


@pytest.mark.parametrize('source,ext',[
    ('image','png'),('webp','jpg'),('png','webp'),('png','pdf'),('xlsx','csv'),
    ('xlsx','json'),('xlsx','xml'),('xlsx','png'),('docx','pdf'),('pptx','pdf'),
    ('pdf','png'),('pdf','txt'),('text','png'),('text','docx'),('text','pptx'),
    ('archive','txt'),('binary','png'),('binary','wav'),('binary','pdf'),
    ('binary','json'),('binary','xml'),('binary','csv'),('binary','xyz'),
    ('png','wav'),('xlsx','wav'),('text','wav'),('wav','png'),('text','zip'),
    ('text','tar'),('text','gz'),('text','bz2'),('text','xz')])
def test_working_conversions(engine,sources,tmp_path,source,ext):
    result=engine.convert(sources[source],tmp_path/f'result.{ext}',Settings(width=480,height=320,audio_seconds=.2,prefer_tts=False))
    assert result.output_size>0 and result.destination_format==ext
    assert not list(tmp_path.glob('.any2any-*'))
    if source=='binary': assert result.classification=='Binary interpretation'
    if ext in {'png','jpg','webp'}:
        with Image.open(result.destination) as image: image.verify()


def test_deterministic_audio(engine,sources,tmp_path):
    settings=Settings(audio_seconds=.2)
    for source in ('png','xlsx','binary'):
        engine.convert(sources[source],tmp_path/'a.wav',settings,overwrite=True)
        engine.convert(sources[source],tmp_path/'b.wav',settings,overwrite=True)
        assert (tmp_path/'a.wav').read_bytes()==(tmp_path/'b.wav').read_bytes()
        import numpy as np
        with wave.open(str(tmp_path/'a.wav')) as audio:
            samples=np.frombuffer(audio.readframes(audio.getnframes()),dtype='<i2')
            assert np.max(np.abs(samples.astype(int)))<=11500


def test_safe_archive_manifest(engine,sources,tmp_path):
    out=tmp_path/'listing.txt'
    engine.convert(sources['archive'],out,Settings(archive_text=True))
    text=out.read_text(encoding='utf-8')
    assert 'Archive text' in text and 'unsafe path; never extracted' in text
    assert not (tmp_path.parent/'unsafe.txt').exists()


def test_bomb_falls_back(engine,tmp_path):
    source=tmp_path/'bomb.zip'
    with zipfile.ZipFile(source,'w',zipfile.ZIP_DEFLATED) as archive:
        archive.writestr('big.txt',b'0'*2_000_000)
    result=engine.convert(source,tmp_path/'safe.txt',Settings(max_archive_bytes=1024))
    assert result.classification=='Binary interpretation'
    assert any('safety limit' in warning for warning in result.warnings)


@pytest.mark.parametrize('header',[b'\x89PNG\r\n\x1a\n',b'%PDF-1.7',b'PK\x03\x04'])
def test_corrupted_inputs(engine,tmp_path,header):
    source=tmp_path/'broken.bin'; source.write_bytes(header+b'broken contents')
    result=engine.convert(source,tmp_path/'result.png',Settings(width=320,height=240))
    assert result.classification=='Binary interpretation' and result.output_size>0
    assert source.read_bytes()==header+b'broken contents'
    assert not list(tmp_path.glob('.any2any-*'))


def test_huge_file_bounded_prefix(engine,tmp_path):
    source=tmp_path/'huge.bin'
    with source.open('wb') as stream:
        stream.write(b'\0'*4096); stream.seek(64*1024*1024-1); stream.write(b'\0')
    output=tmp_path/'output.json'
    result=engine.convert(source,output,Settings(max_read_bytes=1024))
    data=json.loads(output.read_text())
    assert data['original_size']==64*1024*1024 and data['included_bytes']==1024 and data['truncated']
    assert len(output.read_bytes())<3000 and len(data['sha256'])==64
    assert result.warnings


def test_cancellation_preserves_existing_and_cleanup(engine,tmp_path):
    source=tmp_path/'input.bin'; source.write_bytes(b'\0'*3_000_000)
    output=tmp_path/'output.json'; output.write_text('keep me')
    event=threading.Event()
    def progress(amount): event.set()
    with pytest.raises(Cancelled): engine.convert(source,output,Settings(),event,progress=progress,overwrite=True)
    assert output.read_text()=='keep me' and not list(tmp_path.glob('.any2any-*'))
    with pytest.raises(Cancelled): engine.convert(source,tmp_path/'cancelled.json',cancel=event)


def test_overwrite_and_source_protection(engine,sources,tmp_path):
    source=sources['image']; original=source.read_bytes()
    with pytest.raises(ConversionError): engine.convert(source,source,overwrite=True)
    alias=tmp_path/'alias.jpg'
    __import__('os').link(source,alias)
    with pytest.raises(ConversionError): engine.convert(source,alias,overwrite=True)
    output=tmp_path/'existing.png'; output.write_bytes(b'original')
    with pytest.raises(ConversionError): engine.convert(source,output)
    assert output.read_bytes()==b'original' and source.read_bytes()==original


def test_failure_cleanup_and_transaction(engine,sources,tmp_path,monkeypatch):
    import any2any.native as native
    output=tmp_path/'result.png'; output.write_bytes(b'keep')
    original=native.copy
    def broken(ctx,value):
        ctx.destination.write_bytes(b'partial'); raise OSError('simulated disk error')
    monkeypatch.setattr('any2any.writers.binary_writer',broken)
    monkeypatch.setattr('any2any.writers.image_writer',broken)
    with pytest.raises(ConversionError): engine.convert(sources['binary'],output,overwrite=True)
    assert output.read_bytes()==b'keep' and not list(tmp_path.glob('.any2any-*'))


@pytest.mark.parametrize('ext',['../png','a/b','.', '', 'png;whoami'])
def test_extension_sanitization(ext):
    with pytest.raises(ConversionError): normalize_extension(ext)


def test_formula_strings_are_data(engine,tmp_path):
    source=tmp_path/'data.csv'; source.write_text('=1+1,hello\n',encoding='utf-8')
    output=tmp_path/'data.xlsx'; engine.convert(source,output)
    workbook=load_workbook(output)
    assert workbook.active['A1'].data_type=='s' and workbook.active['A1'].value=='=1+1'
    workbook.close()


def test_missing_encoders_use_honest_envelope(engine,sources,tmp_path,monkeypatch):
    import any2any.planner as planner
    monkeypatch.setattr(planner,'locate',lambda name:None)
    result=engine.convert(sources['text'],tmp_path/'result.mp4',Settings(prefer_tts=False))
    assert result.classification=='Binary interpretation'
    assert json.loads(Path(result.destination).read_text())['data_encoding']=='base64'
    assert any('not a valid MP4' in warning for warning in result.warnings)


@pytest.mark.skipif(not locate('ffmpeg'),reason='FFmpeg required')
def test_media_examples_and_subtitles(engine,sources,tmp_path):
    settings=Settings(width=320,height=240,fps=10,audio_seconds=.2,seconds_per_page=.3,max_pages=3,subtitles=True,transition='fade')
    results=[]
    video=tmp_path/'slideshow.mp4'; results.append(engine.convert(sources['pdf'],video,settings))
    assert results[-1].classification=='Semantic conversion'
    info=engine.detector.detect(video)
    assert info.category=='video' and any(stream['codec_type']=='subtitle' for stream in info.details['streams'])
    for ext in ('mov','mkv','avi'):
        results.append(engine.convert(video,tmp_path/f'video.{ext}',settings))
        assert results[-1].classification=='Native conversion'
    results.append(engine.convert(video,tmp_path/'frames.pdf',settings))
    assert results[-1].classification=='Semantic conversion'
    results.append(engine.convert(sources['wav'],tmp_path/'sound.mp3',settings))
    assert results[-1].classification=='Native conversion'
    engine.convert(tmp_path/'sound.mp3',tmp_path/'decoded.wav',settings)
    for view in ('spectrogram','spectrum'):
        result=engine.convert(sources['wav'],tmp_path/f'{view}.png',Settings(audio_view=view,audio_seconds=.2,width=320,height=240))
        assert result.classification=='Semantic conversion'
    gif=tmp_path/'animated.gif'; Image.new('RGB',(64,64),'red').save(gif,save_all=True,append_images=[Image.new('RGB',(64,64),'blue')],duration=100,loop=0)
    result=engine.convert(gif,tmp_path/'gif.mp4',settings)
    assert result.classification=='Native conversion'
    result=engine.convert(sources['xlsx'],tmp_path/'table.mp3',settings)
    assert result.classification=='Semantic conversion'


def test_gui_smoke():
    from tkinterdnd2 import TkinterDnD
    from any2any.gui import App
    root=TkinterDnD.Tk(); root.withdraw()
    try:
        app=App(root); root.update()
        assert app.get_settings().max_pages==50
        assert app.source_entry.winfo_exists()
        app.append_log('GUI smoke test'); root.update()
        assert 'GUI smoke test' in app.log_text.get('1.0','end')
    finally: root.destroy()


def test_external_tool_cancellation(engine,sources,tmp_path,monkeypatch):
    import sys,time
    from any2any.tools import run
    event=threading.Event()
    def slow_writer(ctx,value):
        run(ctx,[sys.executable,'-c','import time; time.sleep(30)'])
    monkeypatch.setattr('any2any.writers.image_writer',slow_writer)
    timer=threading.Timer(.3,event.set); timer.start(); start=time.monotonic()
    try:
        with pytest.raises(Cancelled): engine.convert(sources['image'],tmp_path/'cancel.png',cancel=event)
        assert time.monotonic()-start<6
        assert not (tmp_path/'cancel.png').exists() and not list(tmp_path.glob('.any2any-*'))
    finally: timer.cancel()


def test_retry_semantic_after_native_failure(engine,sources,tmp_path,monkeypatch):
    def failed(ctx,value): raise ConversionError('Simulated unavailable native encoder')
    monkeypatch.setattr('any2any.native.copy',failed)
    result=engine.convert(sources['png'],tmp_path/'copy.png')
    assert result.classification=='Native conversion' and 'Pillow' in result.route
    assert any('next available route' in warning for warning in result.warnings)


def test_media_magic_without_ffprobe(engine,tmp_path,monkeypatch):
    monkeypatch.setattr('any2any.detect.locate',lambda name:None)
    for header,fmt,cat in [(b'\0\0\0\x18ftypisom' + b'\0'*32,'mp4','video'),
                          (b'\0\0\0\x18ftypM4A ' + b'\0'*32,'m4a','audio'),
                          (b'ID3'+b'\0'*32,'mp3','audio'),
                          (b'\x1aE\xdf\xa3'+b'webm'+b'\0'*32,'webm','video')]:
        source=tmp_path/'disguised.bin'; source.write_bytes(header)
        info=engine.detector.detect(source)
        assert info.format==fmt and info.category==cat


def test_planner_bottleneck_preserves_shorter_routes(monkeypatch):
    from any2any.planner import ConversionRegistry,ConversionPlanner
    from any2any.model import FormatInfo
    registry=ConversionRegistry()
    for source,target,name,priority in [('input:binary','a','native decoder',1),
            ('a','b','native adapter',1),('b','common','native join',1),
            ('input:binary','common','semantic shortcut',2),('common','output:xyz','binary writer',4)]:
        registry.register(source,target,name,priority,lambda ctx,value:None)
    monkeypatch.setattr('any2any.planner.default_registry',lambda *args:registry)
    plan=ConversionPlanner().plan(FormatInfo('binary','binary','application/octet-stream','Binary','magic'),'xyz',Settings())
    assert plan.priority==4 and plan.steps[0].name=='semantic shortcut' and len(plan.steps)==2


def test_native_animation_preserves_size_and_timing(engine,tmp_path):
    source=tmp_path/'animation.gif'
    Image.new('RGB',(1400,80),'red').save(source,save_all=True,
        append_images=[Image.new('RGB',(1400,80),'blue')],duration=[100,200],loop=2)
    engine.convert(source,tmp_path/'first.png')
    with Image.open(tmp_path/'first.png') as image: assert image.size==(1400,80)
    result=engine.convert(source,tmp_path/'animation.webp')
    assert result.classification=='Native conversion'
    with Image.open(tmp_path/'animation.webp') as image:
        assert image.size==(1400,80) and image.n_frames==2 and image.info['loop']==2
        image.load(); assert image.info['duration']==100
        image.seek(1); image.load(); assert image.info['duration']==200
