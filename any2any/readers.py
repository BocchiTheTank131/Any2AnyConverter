"""Bounded decoders into reusable intermediate representations."""
import csv
import io
import json
import math
from pathlib import PurePosixPath
import tarfile
import zipfile
import xml.etree.ElementTree as ET
from .model import TextData, TableData, ImageData, DocumentPages, BinaryData, ConversionError
from .tools import locate, run


def small_container(ctx):
    if ctx.source.stat().st_size > ctx.settings.max_read_bytes:
        raise ConversionError('Container exceeds configured read limit. Use binary interpretation or raise the limit.')
    if zipfile.is_zipfile(ctx.source):
        with zipfile.ZipFile(ctx.source) as archive:
            if len(archive.filelist) > ctx.settings.max_archive_entries:
                raise ConversionError('Archive entry limit exceeded.')
            if sum(item.file_size for item in archive.filelist) > ctx.settings.max_archive_bytes:
                raise ConversionError('Decompressed container exceeds safety limit.')
            if any(item.file_size > 1024*1024 and item.file_size / max(1, item.compress_size) > 200
                   for item in archive.filelist):
                raise ConversionError('Suspicious archive compression ratio exceeds safety limit.')


def decode_text(data):
    for encoding in ('utf-8-sig', 'utf-16' if data.startswith((b'\xff\xfe', b'\xfe\xff')) else 'utf-8'):
        try:
            return data.decode(encoding)
        except UnicodeError:
            pass
    try:
        from charset_normalizer import from_bytes
        match = from_bytes(data).best()
        if match is not None:
            return str(match)
    except ImportError:
        pass
    return data.decode('utf-8', 'replace')


def text_reader(ctx, _):
    fmt = ctx.info.format
    if fmt == 'pdf':
        import pymupdf as fitz
        chunks = []
        with fitz.open(ctx.source) as document:
            for page in list(range(min(len(document), ctx.settings.max_pages))):
                ctx.check()
                chunks.append(f'--- Page {page+1} ---\n'+document[page].get_text()[:ctx.settings.max_read_bytes // ctx.settings.max_pages])
            if len(document) > ctx.settings.max_pages: ctx.warn('PDF text limited to maximum page count.')
        return TextData('\n\n'.join(chunks))
    if fmt in {'docx', 'pptx'}:
        small_container(ctx)
        if fmt == 'docx':
            from docx import Document
            document = Document(ctx.source)
            chunks = [p.text for p in document.paragraphs]
            chunks += ['\t'.join(c.text for c in row.cells) for table in document.tables for row in table.rows]
            ctx.warn('DOCX text is extracted; original layout, styles and embedded artwork are not preserved.')
        else:
            from pptx import Presentation
            document = Presentation(ctx.source)
            chunks = []
            for slide in list(document.slides)[:ctx.settings.max_pages]:
                ctx.check()
                chunks.append('\n'.join(shape.text for shape in slide.shapes if shape.has_text_frame))
            ctx.warn('Presentation text is extracted; original layout and artwork are not preserved.')
        return TextData('\n\n'.join(chunks)[:ctx.settings.max_read_bytes])
    return TextData(decode_text(ctx.read()))


def table_reader(ctx, _):
    fmt = ctx.info.format
    if fmt == 'xlsx':
        small_container(ctx)
        from openpyxl import load_workbook
        workbook = load_workbook(ctx.source, read_only=True, data_only=True)
        sheets, count = [], 0
        try:
            for sheet in workbook:
                rows = []
                for row in sheet.iter_rows(values_only=True):
                    ctx.check()
                    remaining = ctx.settings.max_cells - count
                    if remaining <= 0:
                        ctx.warn('Spreadsheet cell limit reached.')
                        break
                    values = list(row[:remaining])
                    rows.append(values)
                    count += len(values)
                sheets.append((sheet.title, rows))
                if count >= ctx.settings.max_cells: break
        finally:
            workbook.close()
        ctx.warn('Spreadsheet conversions use cached formula values; formulas, styles and macros are not retained.')
        return TableData(sheets)
    data = decode_text(ctx.read())
    if fmt == 'json':
        obj = json.loads(data)
        if isinstance(obj,dict) and isinstance(obj.get('sheets'),dict):
            sheets=[]; remaining=ctx.settings.max_cells
            for name,sheet in obj['sheets'].items():
                if not isinstance(sheet,list): raise ConversionError('Invalid structured sheet.')
                bounded=[]
                for row in sheet:
                    ctx.check()
                    if not isinstance(row,list): row=[row]
                    if remaining<=0: ctx.warn('Table cell limit reached.'); break
                    row=row[:remaining]; remaining-=len(row); bounded.append(row)
                sheets.append((str(name),bounded))
            return TableData(sheets or [('Data',[])])
        if isinstance(obj, dict): obj = [obj]
        if not isinstance(obj, list): obj = [[obj]]
        if obj and all(isinstance(row, dict) for row in obj):
            keys = list(dict.fromkeys(k for row in obj for k in row))
            rows = [keys] + [[row.get(k) for k in keys] for row in obj]
        else:
            rows = [row if isinstance(row, list) else [row] for row in obj]
    elif fmt == 'xml':
        if '<!DOCTYPE' in data.upper() or '<!ENTITY' in data.upper():
            raise ConversionError('XML entity/DTD declarations are disabled.')
        root = ET.fromstring(data)
        rows = [['tag', 'text', 'attributes']] + [[node.tag, node.text or '', json.dumps(node.attrib)] for node in root.iter()]
    else:
        rows = []; count=0
        for row in csv.reader(io.StringIO(data), delimiter='\t' if fmt == 'tsv' else ','):
            ctx.check()
            rows.append(row)
            count+=len(row)
            if count >= ctx.settings.max_cells: ctx.warn('Table cell limit reached.'); break
    bounded, count = [], 0
    for row in rows:
        ctx.check()
        if count >= ctx.settings.max_cells:
            ctx.warn('Table cell limit reached.')
            break
        row = row[:ctx.settings.max_cells-count]
        bounded.append(row)
        count += len(row)
    return TableData([('Data', bounded)])


def image_reader(ctx, _):
    from PIL import Image, ImageOps
    from .detect import AUDIO_EXTS,IMAGE_EXTS
    pages = []; durations=[]
    with Image.open(ctx.source) as source:
        if source.width * source.height > 40_000_000:
            raise ConversionError('Decoded image exceeds the 40 megapixel safety limit.')
        total_pixels=0
        total_frames=getattr(source,'n_frames',1)
        extension=ctx.destination.suffix.lower()[1:]
        maximum=1 if extension in AUDIO_EXTS else ctx.settings.max_pages
        loop=source.info.get('loop',0)
        for i in range(min(total_frames,maximum)):
            ctx.check()
            source.seek(i)
            if source.width * source.height > 40_000_000:
                raise ConversionError('Decoded image exceeds the 40 megapixel safety limit.')
            image = ImageOps.exif_transpose(source).convert('RGBA')
            # Bound cumulative page memory; resize only multi-frame semantic intermediates.
            if total_frames > 1 and extension not in IMAGE_EXTS|{'pdf'}:
                image.thumbnail((ctx.settings.width, ctx.settings.height))
            total_pixels+=image.width*image.height
            if total_pixels>40_000_000:
                ctx.warn('Image sequence capped at 40 million total rendered pixels.'); break
            pages.append(image)
            durations.append(int(source.info.get('duration',ctx.settings.seconds_per_page*1000)))
        if total_frames>maximum: ctx.warn('Image frames limited to first page for audio or the configured maximum page count.')
    return ImageData(pages,durations=durations if total_frames>1 else [],loop=loop)


def pdf_reader(ctx, _):
    import pymupdf as fitz
    from PIL import Image
    pages, captions = [], []
    with fitz.open(ctx.source) as document:
        for index in range(min(len(document), ctx.settings.max_pages)):
            ctx.check()
            page = document[index]
            max_count=min(len(document),ctx.settings.max_pages)
            budget_scale=math.sqrt(40_000_000/max(1,max_count)/(page.rect.width*page.rect.height))
            scale = min(ctx.settings.width / page.rect.width, ctx.settings.height / page.rect.height,budget_scale)
            pix = page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False)
            pages.append(Image.frombytes('RGB', (pix.width, pix.height), pix.samples))
            captions.append(page.get_text()[:3000] if ctx.settings.subtitles else f'Page {index+1}')
            ctx.progress(.15 + .35*(index+1)/min(len(document), ctx.settings.max_pages))
        if len(document) > ctx.settings.max_pages: ctx.warn('PDF rendering limited to maximum page count.')
    return DocumentPages(pages, captions)


def video_reader(ctx, _):
    from PIL import Image
    ffmpeg = locate('ffmpeg')
    if not ffmpeg: raise ConversionError('FFmpeg unavailable.')
    settings = ctx.settings
    max_width=min(settings.width,int(math.sqrt(40_000_000/settings.max_pages*settings.width/settings.height)))
    max_height=min(settings.height,int(40_000_000/settings.max_pages/max(1,max_width)))
    if settings.frame_mode == 'every':
        selection = 'select=1'
    elif settings.frame_mode == 'frames':
        selection = f'select=not(mod(n\\,{max(1, int(settings.frame_interval))}))'
    else:
        selection = f'fps=1/{settings.frame_interval}:start_time=0:round=up:eof_action=pass'
    ctx.warn(f'Video extraction is capped at {settings.max_pages} pages/frames.')
    filter_ = selection + f',scale={max_width}:{max_height}:force_original_aspect_ratio=decrease,showinfo'
    run(ctx, [ffmpeg, '-nostdin', '-v', 'info', '-i', ctx.source, '-vf', filter_, '-fps_mode', 'vfr',
              '-frames:v', str(settings.max_pages), ctx.temporary / 'frame-%05d.png'])
    if not list(ctx.temporary.glob('frame-*.png')):
        ctx.warn('Video is shorter than the sampling interval; using its first decodable frame.')
        run(ctx,[ffmpeg,'-nostdin','-v','info','-i',ctx.source,'-vf',
                 f'scale={max_width}:{max_height}:force_original_aspect_ratio=decrease,showinfo',
                 '-frames:v','1',ctx.temporary/'frame-00001.png'])
    import re
    log = (ctx.temporary / 'external.log').read_text(errors='replace')
    # showinfo follows sampling; output PTS for fps filter reflects selected sampling timeline.
    times = re.findall(r'pts_time:([\d.]+)', log)
    pages, captions = [], []
    for index, path in enumerate(sorted(ctx.temporary.glob('frame-*.png'))):
        ctx.check()
        with Image.open(path) as image: pages.append(image.copy())
        captions.append(f'{float(times[index]):.2f} s' if index < len(times) else f'Frame {index+1}')
    if not pages: raise ConversionError('Video has no decodable frames.')
    return ImageData(pages, captions if settings.timestamps else [])


def archive_reader(ctx, _):
    small_container(ctx)
    lines = [f'Archive: {ctx.source.name}', f'Format: {ctx.info.format}', f'Size: {ctx.source.stat().st_size} bytes', '']
    entries = 0
    text_bytes = 0
    def add(name, size, reader=None):
        nonlocal entries, text_bytes
        ctx.check()
        entries += 1
        if entries > ctx.settings.max_archive_entries: raise ConversionError('Archive entry limit exceeded.')
        safe = not (PurePosixPath(name.replace('\\', '/')).is_absolute() or
                    '..' in PurePosixPath(name.replace('\\', '/')).parts or ':' in name)
        lines.append(f'{size:>12}  {name}' + (' [unsafe path; never extracted]' if not safe else ''))
        if ctx.settings.archive_text and safe and reader and size <= 65536 and text_bytes+size <= ctx.settings.max_archive_bytes:
            if PurePosixPath(name).suffix.lower() in {'.txt', '.md', '.csv', '.json', '.xml', '.log'}:
                text_bytes += size
                lines.append(decode_text(reader())[:65536])
    fmt = ctx.info.format
    if fmt == 'zip':
        with zipfile.ZipFile(ctx.source) as archive:
            for item in archive.infolist():
                add(item.filename, item.file_size, lambda item=item: archive.read(item))
    elif fmt == '7z':
        import py7zr
        with py7zr.SevenZipFile(ctx.source) as archive:
            for item in archive.list(): add(item.filename, item.uncompressed or 0)
        if ctx.settings.archive_text: ctx.warn('7Z lists metadata only; contents are never extracted.')
    elif fmt == 'tar' or tarfile.is_tarfile(ctx.source):
        total = 0
        with tarfile.open(ctx.source, 'r|*') as archive:
            for item in archive:
                total += item.size
                if total > ctx.settings.max_archive_bytes: raise ConversionError('Archive expanded-size limit exceeded.')
                def read_member(item=item):
                    with archive.extractfile(item) as stream: return stream.read(65536)
                add(item.name, item.size, read_member if item.isfile() else None)
    else:
        import gzip, bz2, lzma
        opener = {'gz': gzip.open, 'bz2': bz2.open, 'xz': lzma.open}[fmt]
        with opener(ctx.source, 'rb') as stream:
            content = stream.read(min(ctx.settings.max_archive_bytes, ctx.settings.max_read_bytes)+1)
        if len(content) > min(ctx.settings.max_archive_bytes, ctx.settings.max_read_bytes):
            raise ConversionError('Compressed stream exceeds expanded-size limit.')
        add(ctx.source.stem, len(content), lambda: content)
        if ctx.settings.archive_text: lines.append(decode_text(content))
    return TextData('\n'.join(lines)[:ctx.settings.max_read_bytes])


def binary_reader(ctx, _):
    import hashlib
    digest = hashlib.sha256()
    total = ctx.source.stat().st_size
    count = 0
    with ctx.source.open('rb') as stream:
        while chunk := stream.read(1024*1024):
            ctx.check()
            digest.update(chunk)
            count += len(chunk)
            ctx.progress(.1 + .15*count/max(1,total))
    data = ctx.read()
    return BinaryData(data, {'conversion_type': 'binary_fallback', 'original_name': ctx.source.name,
                           'original_format': ctx.info.description, 'original_size': total,
                           'sha256': digest.hexdigest(), 'included_bytes': len(data), 'truncated': len(data)<total})
