"""Packaged-runtime smoke test; writes a report without requiring a console."""
from pathlib import Path
from dataclasses import asdict
import json
import traceback
from .engine import ConversionEngine
from .model import Settings
from .tools import capabilities


def run(folder):
    folder=Path(folder).resolve(); folder.mkdir(parents=True,exist_ok=True)
    report={'capabilities':capabilities(),'results':[]}
    try:
        from tkinterdnd2 import TkinterDnD
        from .gui import App
        root=TkinterDnD.Tk(); root.withdraw()
        try:
            app=App(root); root.update(); app.get_settings()
            report['gui']='Tk + drag-and-drop initialized successfully'
        finally: root.destroy()
        from PIL import Image
        import pymupdf as fitz
        from openpyxl import Workbook
        from docx import Document
        image=folder/'input.png'; Image.new('RGB',(96,64),'#38bdf8').save(image)
        document=fitz.open()
        for title in ('First page','Second page'):
            page=document.new_page(width=300,height=200); page.insert_text((20,40),title)
        document.save(folder/'input.pdf'); document.close()
        workbook=Workbook(); workbook.active.append(['one','two']); workbook.active.append([12,25]); workbook.save(folder/'input.xlsx')
        word=Document(); word.add_paragraph('Packaged document test'); word.save(folder/'input.docx')
        (folder/'input.bin').write_bytes(bytes(range(256))*4)
        settings=Settings(width=320,height=240,seconds_per_page=.5,max_pages=3,frame_interval=.2,
                          audio_seconds=.2,subtitles=True,transition='fade',prefer_tts=False)
        engine=ConversionEngine()
        for source,dest in [('input.png','image.jpg'),('input.png','image.wav'),('input.xlsx','table.mp3'),
                            ('input.docx','word.pdf'),('input.pdf','slideshow.mp4'),('slideshow.mp4','frames.pdf'),
                            ('input.bin','bytes.xyz'),('image.wav','audio-report.txt'),('slideshow.mp4','video-report.txt'),
                            ('input.png','image-report.json'),('input.xlsx','table-report.txt')]:
            result=engine.convert(folder/source,folder/dest,settings,overwrite=True)
            if result.classification=='Binary interpretation' and source!='input.bin':
                raise RuntimeError(f'Unexpected fallback: {result.warnings}')
            report['results'].append(asdict(result))
        from dataclasses import replace
        (folder/'input.txt').write_text('Meaningful scrolling text\nSecond line',encoding='utf-8')
        report['results'].append(asdict(engine.convert(folder/'input.txt',folder/'scroll.mp4',replace(settings,mode='scroll'),overwrite=True)))
        report['results'].append(asdict(engine.convert(folder/'slideshow.mp4',folder/'contact.png',replace(settings,mode='contact'),overwrite=True)))
        from .batch import BatchConverter
        batch=BatchConverter(engine)
        items=batch.plan([folder/'input.png',folder/'input.xlsx'],folder,'webp',False)
        batch.convert(items,settings,overwrite=True)
        if any(item.status!='complete' for item in items): raise RuntimeError('Packaged batch failed: '+str([item.error for item in items]))
        report['batch']=[asdict(item) for item in items]
        report['success']=True
    except Exception:
        report['success']=False; report['error']=traceback.format_exc()
    (folder/'report.json').write_text(json.dumps(report,indent=2),encoding='utf-8')
    return report['success']
