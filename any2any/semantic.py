"""Local extraction and analysis bridges. Optional recognition never downloads models."""
from dataclasses import replace
from html.parser import HTMLParser
import importlib.util
import json
import math
import os
from pathlib import Path
import re
import sys
import wave
import zipfile
import xml.etree.ElementTree as ET
from .model import (AudioData, VideoData, TextData, MetadataData, SubtitleData,
                    ImageData, DocumentPages, ConversionError)
from .tools import locate, run
from . import readers


def optional(ctx, title, operation):
    try:
        ctx.check()
        return operation()
    except Exception as error:
        ctx.check()  # Cancellation must never become a failed optional feature.
        ctx.warn(f'{title} unavailable/failed: {str(error)[:500]}')
        return None


def audio_source(ctx, _): return AudioData(ctx.source)
def video_source(ctx, _): return VideoData(ctx.source)


def probe(ctx):
    cached = ctx.info.details.get('probe')
    if cached: return cached
    tool = locate('ffprobe')
    if tool:
        run(ctx, [tool, '-protocol_whitelist', 'file,pipe', '-v', 'error', '-show_format',
                  '-show_streams', '-of', 'json', ctx.source])
        with (ctx.temporary / 'external.log').open('rb') as stream:
            result = json.loads(stream.read(ctx.settings.max_read_bytes))
        ctx.info.details['probe'] = result
        return result
    if ctx.info.details.get('streams'):
        return {'format':{'duration':ctx.info.details.get('duration')},'streams':ctx.info.details['streams']}
    tool=locate('ffmpeg')
    if not tool: return {'format':{},'streams':[]}
    run(ctx,[tool,'-nostdin','-v','info','-i',ctx.source,'-t','0.01','-f','null','-'])
    log=(ctx.temporary/'external.log').read_text(errors='replace')
    # Keep input header only; output headers would describe the analysis decoder's format.
    header=log.split('Stream mapping:',1)[0]
    fields={'format_name':ctx.info.format,'size':str(ctx.source.stat().st_size)}
    match=re.search(r'Duration: (\d+):(\d+):(\d+(?:\.\d+)?)',header)
    if match: fields['duration']=sum(float(v)*factor for v,factor in zip(match.groups(),(3600,60,1)))
    match=re.search(r'bitrate:\s*(\d+) kb/s',header)
    if match: fields['bit_rate']=int(match.group(1))*1000
    streams=[]
    for line in header.splitlines():
        match=re.search(r'Stream #0:(\d+).*?: (Audio|Video|Subtitle): ([\w-]+)',line)
        if not match: continue
        stream={'index':int(match.group(1)),'codec_type':match.group(2).lower(),'codec_name':match.group(3)}
        rate=re.search(r'(\d+) Hz',line)
        if rate: stream['sample_rate']=int(rate.group(1))
        if 'stereo' in line: stream['channels']=2
        elif 'mono' in line: stream['channels']=1
        dims=re.search(r'\b(\d{2,5})x(\d{2,5})\b',line)
        if dims: stream.update(width=int(dims.group(1)),height=int(dims.group(2)))
        streams.append(stream)
    result={'format':fields,'streams':streams}
    ctx.info.details['probe']=result
    ctx.warn('ffprobe unavailable; media metadata was parsed from the bundled FFmpeg header and may be less complete.')
    return result


def media_metadata(ctx):
    data = probe(ctx)
    format_ = data.get('format', {})
    fields = {key: format_[key] for key in ('format_name', 'duration', 'size', 'bit_rate', 'tags') if format_.get(key) is not None}
    fields['streams'] = [{k: v for k, v in stream.items() if k in {
        'index', 'codec_type', 'codec_name', 'channels', 'sample_rate', 'bit_rate',
        'width', 'height', 'avg_frame_rate', 'duration', 'tags'}} for stream in data.get('streams', [])]
    return MetadataData(ctx.info.description, fields)


def metadata_text(ctx, value):
    return TextData(value.title + '\n' + json.dumps(value.fields, ensure_ascii=False, indent=2, default=str))


def decoded_audio(ctx, channels=1):
    tool = locate('ffmpeg')
    path = ctx.temporary / f'analysis-{channels}.wav'
    if tool:
        run(ctx, [tool, '-nostdin', '-v', 'error', '-i', ctx.source, '-vn',
                  '-t', str(ctx.settings.audio_seconds), '-ac', str(channels),
                  '-ar', str(ctx.settings.sample_rate), '-c:a', 'pcm_s16le', path])
    elif ctx.info.format == 'wav':
        path = ctx.source
    else: raise ConversionError('Audio decoding requires FFmpeg.')
    return path


def audio_analysis(ctx):
    import numpy as np
    path = decoded_audio(ctx)
    with wave.open(str(path), 'rb') as stream:
        if stream.getsampwidth() != 2: raise ConversionError('Analysis requires 16-bit PCM.')
        rate, channels = stream.getframerate(), stream.getnchannels()
        data = stream.readframes(int(ctx.settings.audio_seconds * rate))
    samples = np.frombuffer(data, dtype='<i2').astype(float) / 32768
    if not len(samples): raise ConversionError('No audio samples to analyze.')
    samples = samples.reshape(-1, channels).mean(axis=1)
    rms = math.sqrt(float(np.mean(samples ** 2)))
    peak = float(np.max(np.abs(samples)))
    windows = max(1, int(rate * .1)); silence = []; start = None
    for index in range(0, len(samples), windows):
        ctx.check()
        quiet = math.sqrt(float(np.mean(samples[index:index+windows] ** 2))) < .01
        if quiet and start is None: start = index/rate
        if not quiet and start is not None:
            if index/rate-start >= .3: silence.append((round(start,2), round(index/rate,2)))
            start = None
    if start is not None and len(samples)/rate-start >= .3:
        silence.append((round(start,2), round(len(samples)/rate,2)))
    return MetadataData('Audio analysis (bounded decoded prefix)', {
        'analyzed_seconds': round(len(samples)/rate,3), 'analysis_sample_rate': rate,
        'rms_level_dbfs': round(20*math.log10(max(rms,1e-12)),2),
        'peak_level_dbfs': round(20*math.log10(max(peak,1e-12)),2),
        'silence_sections_seconds': silence[:1000], 'silence_threshold_dbfs': -40,
        'note': 'RMS level is a loudness estimate, not calibrated LUFS. Silence windows are 0.1 s; sections >= 0.3 s.'})


def transcription_backend():
    model = os.environ.get('ANY2ANY_WHISPER_MODEL', '')
    if not model or not Path(model).exists(): return None
    if Path(model).is_dir() and importlib.util.find_spec('faster_whisper'): return 'faster_whisper'
    if Path(model).is_file() and Path(model).suffix == '.pt' and importlib.util.find_spec('whisper'): return 'whisper'
    if Path(model).is_file() and locate('whisper-cli'): return 'whisper-cli'
    return None


def transcribe(ctx):
    if not ctx.settings.transcribe: return None
    backend = transcription_backend()
    if not backend:
        ctx.warn('Offline transcription requested but no local model/backend is configured. Set ANY2ANY_WHISPER_MODEL.')
        return None
    path = decoded_audio(ctx)
    output = ctx.temporary / 'transcript.txt'
    model = os.environ['ANY2ANY_WHISPER_MODEL']
    if backend == 'whisper-cli':
        prefix = ctx.temporary / 'transcript'
        run(ctx, [locate('whisper-cli'), '-m', model, '-f', path, '-otxt', '-of', prefix])
    else:
        args = [sys.executable, '--transcribe-worker'] if getattr(sys, 'frozen', False) else [sys.executable, '-m', 'any2any.transcribe_worker']
        run(ctx, [*args, backend, model, path, output])
    with output.open('rb') as stream: text = readers.decode_text(stream.read(ctx.settings.max_read_bytes))
    ctx.warn(f'Offline transcription is limited to the first {ctx.settings.audio_seconds:g} seconds; recognition can contain errors.')
    return TextData(text)


def subtitles(ctx):
    data = probe(ctx)
    streams = [s for s in data.get('streams', []) if s.get('codec_type') == 'subtitle']
    parts = []
    for stream in streams[:8]:
        ctx.check()
        path = ctx.temporary / f'subtitle-{stream["index"]}.srt'
        def extract():
            run(ctx, [locate('ffmpeg'), '-nostdin', '-v', 'error', '-i', ctx.source,
                      '-map', f'0:{stream["index"]}', '-t', str(ctx.settings.audio_seconds), '-c:s', 'srt', path])
            with path.open('rb') as source: return readers.decode_text(source.read(ctx.settings.max_read_bytes))
        text = optional(ctx, 'Embedded subtitle extraction', extract)
        if text and text.strip(): parts.append(text)
    return SubtitleData('\n\n'.join(parts))


def ocr(ctx, image, name='ocr'):
    if not ctx.settings.ocr: return None
    tool = locate('tesseract')
    if not tool:
        ctx.log('OCR unavailable: install Tesseract or set ANY2ANY_TESSERACT.')
        return None
    path = ctx.temporary / (name+'.png'); output = ctx.temporary / name
    image = image.copy(); image.thumbnail((2000,2000)); image.convert('RGB').save(path)
    run(ctx, [tool, path, output, '--psm', '3'])
    with output.with_suffix('.txt').open('rb') as stream:
        return readers.decode_text(stream.read(ctx.settings.max_read_bytes))


def image_text(ctx, _):
    from PIL import Image, ExifTags
    import numpy as np
    with Image.open(ctx.source) as image:
        if image.width*image.height > 40_000_000: raise ConversionError('Image exceeds analysis pixel limit.')
        fields = {'format': image.format, 'dimensions': image.size, 'mode': image.mode,
                  'frames': getattr(image, 'n_frames', 1)}
        fields['metadata'] = {str(k): str(v)[:2000] for k,v in image.info.items() if k not in {'icc_profile', 'exif'}}
        fields['EXIF'] = {ExifTags.TAGS.get(k,str(k)):str(v)[:2000] for k,v in image.getexif().items()}
        parts = [metadata_text(ctx, MetadataData('Image metadata / EXIF',fields)).text]
        text = optional(ctx, 'Image OCR', lambda: ocr(ctx,image))
        if text and text.strip(): parts.append('OCR (may contain recognition errors):\n'+text)
        # A representation remains useful when there is no text or OCR backend.
        small = image.convert('RGB'); small.thumbnail((512,512))
        rgb = np.asarray(small,dtype=float)
        stats = {'mean_RGB': rgb.mean(axis=(0,1)).round(2).tolist(), 'min_RGB':rgb.min(axis=(0,1)).tolist(),
                 'max_RGB':rgb.max(axis=(0,1)).tolist(), 'brightness_stddev': round(float(rgb.mean(axis=2).std()),2)}
        ascii_image = image.convert('L'); ascii_image.thumbnail((80,40))
        pixels = np.asarray(ascii_image); ramp = ' .:-=+*#%@'
        art = '\n'.join(''.join(ramp[min(9,int(pixel)*10//256)] for pixel in row) for row in pixels)
        parts.append('ASCII-art interpretation:\n'+art)
        parts.append(metadata_text(ctx,MetadataData('Pixel statistics (bounded thumbnail)',stats)).text)
    return TextData('\n\n'.join(parts)[:ctx.settings.max_read_bytes])


def audio_text(ctx, _):
    parts = []
    metadata = optional(ctx,'Audio metadata',lambda:media_metadata(ctx))
    if metadata and (metadata.fields.get('streams') or len(metadata.fields)>1): parts.append(metadata_text(ctx,metadata).text)
    transcript = optional(ctx,'Offline speech recognition',lambda:transcribe(ctx))
    if transcript and transcript.text.strip(): parts.append('Transcription:\n'+transcript.text)
    analysis = optional(ctx,'Audio analysis',lambda:audio_analysis(ctx))
    if analysis: parts.append(metadata_text(ctx,analysis).text)
    if not parts: raise ConversionError('No media metadata, transcript or decoded audio analysis is available.')
    return TextData('\n\n'.join(parts)[:ctx.settings.max_read_bytes])


def video_text(ctx, _):
    parts = []
    extracted = optional(ctx,'Video subtitles',lambda:subtitles(ctx))
    if extracted and extracted.text.strip(): parts.append('Embedded subtitles:\n'+extracted.text)
    metadata = optional(ctx,'Video metadata',lambda:media_metadata(ctx))
    if metadata and (metadata.fields.get('streams') or len(metadata.fields)>1): parts.append(metadata_text(ctx,metadata).text)
    transcript = optional(ctx,'Video speech transcription',lambda:transcribe(ctx))
    if transcript and transcript.text.strip(): parts.append('Speech transcription:\n'+transcript.text)
    if ctx.settings.ocr and locate('tesseract'):
        def frame_ocr():
            from .rendering import video_frames
            local = replace(ctx, settings=replace(ctx.settings,max_pages=min(5,ctx.settings.max_pages),frame_mode='even',video_view='sampling'))
            frames = video_frames(local,None)
            texts = []
            for index,image in enumerate(frames.pages):
                text = optional(ctx,'Frame OCR',lambda image=image,index=index:ocr(ctx,image,f'frame-ocr-{index}'))
                if text and text.strip(): texts.append(f'Frame {index+1}:\n{text}')
            return '\n\n'.join(texts)
        text = optional(ctx,'Sampled frame OCR',frame_ocr)
        if text: parts.append('Sampled frame OCR:\n'+text)
    analysis = optional(ctx,'Video soundtrack analysis',lambda:audio_analysis(ctx))
    if analysis: parts.append(metadata_text(ctx,analysis).text)
    if not parts:
        # Validate one decoded frame before claiming a useful video summary without ffprobe.
        from .rendering import video_frames
        frames = video_frames(replace(ctx,settings=replace(ctx.settings,video_view='first')),None)
        parts.append(f'Decoded video analysis: first frame {frames.pages[0].width} × {frames.pages[0].height} pixels.\n'
                     'Detailed duration/codec metadata requires ffprobe. No subtitles or speech were extracted.')
    return TextData('\n\n'.join(parts)[:ctx.settings.max_read_bytes])


class ReadableHTML(HTMLParser):
    def __init__(self): super().__init__(convert_charrefs=True); self.parts=[]; self.skip=0
    def handle_starttag(self,tag,attrs):
        if tag in {'script','style'}: self.skip+=1
        if tag in {'p','div','br','tr','li','h1','h2','h3','section'} and not self.skip: self.parts.append('\n')
    def handle_endtag(self,tag):
        if tag in {'script','style'} and self.skip: self.skip-=1
        if tag in {'p','div','tr','li','h1','h2','h3'} and not self.skip: self.parts.append('\n')
    def handle_data(self,data):
        if not self.skip: self.parts.append(data)


def document_text(ctx, _):
    fmt = ctx.info.format
    if fmt in {'html','htm'}:
        parser=ReadableHTML(); parser.feed(readers.decode_text(ctx.read()))
        return TextData(re.sub(r'\n[ \t]*\n+', '\n\n',''.join(parser.parts)).strip())
    if fmt=='rtf':
        from striprtf.striprtf import rtf_to_text
        return TextData(rtf_to_text(readers.decode_text(ctx.read()),errors='replace'))
    if fmt in {'odt','odp'}:
        root=office_xml(ctx)
        texts=[]
        for node in root.iter():
            ctx.check()
            if node.tag.rsplit('}',1)[-1] in {'p','h'}: texts.append(''.join(node.itertext()))
        return TextData('\n'.join(texts)[:ctx.settings.max_read_bytes])
    return readers.text_reader(ctx,None)


def office_xml(ctx):
    readers.small_container(ctx)
    with zipfile.ZipFile(ctx.source) as archive: data=archive.read('content.xml')
    if b'<!DOCTYPE' in data.upper() or b'<!ENTITY' in data.upper(): raise ConversionError('Office XML entities/DTD are disabled.')
    return ET.fromstring(data)


def ods_table(ctx,_):
    from .model import TableData
    root=office_xml(ctx); sheets=[]; count=0
    ns='{urn:oasis:names:tc:opendocument:xmlns:table:1.0}'
    for table in root.iter(ns+'table'):
        rows=[]
        for row in table.findall(ns+'table-row'):
            ctx.check(); values=[]
            for cell in row:
                if not cell.tag.endswith('table-cell'): continue
                repeat=min(int(cell.get(ns+'number-columns-repeated','1')),ctx.settings.max_cells-count-len(values))
                value='\n'.join(''.join(p.itertext()) for p in cell if p.tag.rsplit('}',1)[-1]=='p')
                values.extend([value]*max(0,repeat))
            repeats=min(int(row.get(ns+'number-rows-repeated','1')),max(1,(ctx.settings.max_cells-count)//max(1,len(values))))
            for _ in range(repeats):
                if count>=ctx.settings.max_cells: break
                rows.append(values[:ctx.settings.max_cells-count]); count+=max(1,len(rows[-1]))
            if count>=ctx.settings.max_cells: ctx.warn('ODS cell limit reached.'); break
        sheets.append((table.get(ns+'name','Sheet'),rows))
        if count>=ctx.settings.max_cells: break
    return TableData(sheets or [('Sheet',[])])


def presentation_pages(ctx, _):
    from PIL import Image, ImageDraw
    from .bridges import font
    if locate('soffice'):
        from .native import office
        target=ctx.temporary/'presentation.pdf'
        native_ctx=replace(ctx,destination=target)
        result=optional(ctx,'LibreOffice slide rendering',lambda:office(native_ctx,None))
        if target.exists(): return readers.pdf_reader(replace(ctx,source=target),None)
    if ctx.info.format!='pptx':
        from .rendering import text_pages
        ctx.warn('ODP slideshow uses extracted slide text because LibreOffice is unavailable.')
        return text_pages(ctx,document_text(ctx,None))
    readers.small_container(ctx)
    from pptx import Presentation
    presentation=Presentation(ctx.source)
    scale=min(ctx.settings.width/presentation.slide_width,ctx.settings.height/presentation.slide_height)
    w=max(1,int(presentation.slide_width*scale)); h=max(1,int(presentation.slide_height*scale))
    pages=[]
    for slide in list(presentation.slides)[:ctx.settings.max_pages]:
        ctx.check()
        if (len(pages)+1)*w*h>40_000_000: ctx.warn('Slide rendering pixel budget reached.'); break
        image=Image.new('RGB',(w,h),'white'); draw=ImageDraw.Draw(image)
        for shape in slide.shapes:
            x,y=int(shape.left*scale),int(shape.top*scale)
            sw,sh=max(1,int(shape.width*scale)),max(1,int(shape.height*scale))
            if shape.shape_type==13:
                import io
                def picture():
                    with Image.open(io.BytesIO(shape.image.blob)) as source:
                        if source.width*source.height>40_000_000: raise ConversionError('Slide image exceeds pixel limit.')
                        source.thumbnail((sw,sh)); image.paste(source.convert('RGB'),(x,y))
                optional(ctx,'Slide picture',picture)
            elif shape.has_text_frame:
                import textwrap
                line_y=y
                for paragraph in shape.text_frame.paragraphs:
                    point=next((run.font.size.pt for run in paragraph.runs if run.font.size),18)
                    size=max(9,min(64,int(point*12700*scale)))
                    for line in textwrap.wrap(paragraph.text,max(1,int(sw/max(5,size*.6)))):
                        if line_y+size>y+sh: break
                        draw.text((x,line_y),line,font=font(size),fill='#172554'); line_y+=size+3
        pages.append(image)
    ctx.warn('Built-in slide rendering preserves slide aspect ratio, text and pictures; complex shapes, charts, themes and animations may differ. Install LibreOffice for faithful rendering.')
    return DocumentPages(pages)
